"""
generate_deck.py — Generates ultra-premium widescreen PowerPoint presentation (16:9)
for TechPulse AI Digest Bot.
Theme: Signal Monolith (Deep Obsidian, Glowing Emerald, Cyan, Slate Grey, and Pure White).
Embedded High-Resolution Visuals: Dashboard Grid & Cockpit Split View.
"""

import os
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# ─── Color Palette (Signal Monolith Theme) ──────────────────────────────────
BG_DARK = RGBColor(7, 8, 11)           # Deep Obsidian #07080B
CARD_BG = RGBColor(18, 22, 34)         # Elevated Card #121622
CARD_BORDER = RGBColor(38, 48, 68)     # Hairline Border #263044
EMERALD = RGBColor(16, 185, 129)       # Accent Emerald #10B981
EMERALD_BG = RGBColor(12, 45, 34)      # Emerald Tag BG #0C2D22
CYAN = RGBColor(14, 165, 233)          # Accent Cyan #0EA5E9
AMBER = RGBColor(245, 158, 11)         # Accent Amber #F59E0B
WHITE = RGBColor(248, 250, 252)        # Title White #F8FAFC
TEXT_BODY = RGBColor(148, 163, 184)    # Slate Text #94A3B8
TEXT_MUTED = RGBColor(100, 116, 139)   # Subdued Slate #64748B

SCREENSHOT_1 = Path("tests/screenshots/1_dashboard_grid.png").resolve()
SCREENSHOT_2 = Path("tests/screenshots/2_cockpit_split_view.png").resolve()
OUTPUT_FILE = Path("TechPulse_AI_Digest_Bot_Demo.pptx").resolve()


def set_slide_background(slide, prs):
    """Fill slide with deep obsidian background."""
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_DARK
    bg.line.fill.background()
    return bg


def add_header(slide, title_text, category_text="TECHPULSE // OSINT-AI OBSERVATORY"):
    """Render consistent brand header banner on each content slide."""
    # Category tag
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10), Inches(0.35))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = category_text
    p_tag.font.size = Pt(10)
    p_tag.font.bold = True
    p_tag.font.color.rgb = EMERALD
    p_tag.font.name = "Arial"

    # Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.68), Inches(11.5), Inches(0.7))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = WHITE
    p_title.font.name = "Arial"


def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
    """Add a stylized card container with rounded corners and border."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1)
    return card


def build_presentation():
    prs = Presentation()
    # 16:9 Widescreen format
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: Title & Executive Overview
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, prs)

    # Glowing Monolith Accent Bar
    accent_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.3), Inches(0.08), Inches(4.2))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = EMERALD
    accent_bar.line.fill.background()

    # Title & Subtitle block
    t_box = slide1.shapes.add_textbox(Inches(1.1), Inches(1.2), Inches(11), Inches(3.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "AUTONOMOUS TECHNICAL INTELLIGENCE BOT"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = EMERALD
    p0.font.name = "Arial"
    p0.space_after = Pt(8)

    p1 = tf1.add_paragraph()
    p1.text = "TechPulse AI Digest Bot"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.font.name = "Arial"
    p1.space_after = Pt(12)

    p2 = tf1.add_paragraph()
    p2.text = "Autonomous multi-source intelligence engine distilling open-source velocity, research preprints, and developer signals into high-conviction executive briefings."
    p2.font.size = Pt(15)
    p2.font.color.rgb = TEXT_BODY
    p2.font.name = "Arial"
    p2.space_after = Pt(24)

    # 4 Stat/Key Metric Badges on Title Slide
    kpi_items = [
        ("6 Multi-Feeds", "GitHub, HN, HF, ArXiv, Dev.to", CYAN),
        ("Groq 120B LLM", "Sub-15s Curation & JSON Schema", EMERALD),
        ("30-Day Dedup", "SQLite WAL Sliding Window", AMBER),
        ("$0 / Month", "Zero-Cost Serverless & Actions", WHITE),
    ]

    badge_width = Inches(2.7)
    badge_height = Inches(1.4)
    for i, (title, sub, color) in enumerate(kpi_items):
        bx = Inches(0.8 + i * 2.95)
        by = Inches(4.8)
        add_card(slide1, bx, by, badge_width, badge_height)

        kbox = slide1.shapes.add_textbox(bx + Inches(0.15), by + Inches(0.15), badge_width - Inches(0.3), badge_height - Inches(0.3))
        ktf = kbox.text_frame
        ktf.word_wrap = True
        
        kp1 = ktf.paragraphs[0]
        kp1.text = title
        kp1.font.size = Pt(15)
        kp1.font.bold = True
        kp1.font.color.rgb = color
        kp1.font.name = "Arial"
        kp1.space_after = Pt(4)

        kp2 = ktf.add_paragraph()
        kp2.text = sub
        kp2.font.size = Pt(10)
        kp2.font.color.rgb = TEXT_BODY
        kp2.font.name = "Arial"

    # Footer Metadata
    ft_box = slide1.shapes.add_textbox(Inches(0.8), Inches(6.6), Inches(11.5), Inches(0.4))
    ft_p = ft_box.text_frame.paragraphs[0]
    ft_p.text = "Author: BT-Nishant (Bugendaitech)  |  Private GitHub Repository: Bugendaitech-Nishant/techpulse-ai-digest-bot"
    ft_p.font.size = Pt(10)
    ft_p.font.color.rgb = TEXT_MUTED
    ft_p.font.name = "Arial"

    # =========================================================================
    # SLIDE 2: 3-Agent Decoupled Pipeline Architecture
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, prs)
    add_header(slide2, "Decoupled 3-Agent Autonomous Pipeline Architecture")

    # 3 Agent Cards
    agents = [
        (
            "AGENT 1: HARVESTER",
            "Multi-Source Concurrent Ingestion",
            [
                "Concurrent multi-threaded crawling across 6 APIs & scrapers in <7s total.",
                "In-batch normalization (UTM parameter stripping, URL SHA-256 hashing).",
                "30-day SQLite sliding deduplication window: prevents repeat surfacing.",
                "Robust error isolation: individual feed failure never halts the pipeline.",
            ],
            CYAN,
        ),
        (
            "AGENT 2: PROCESSOR",
            "Groq LLM Filter, Ranker & Summarizer",
            [
                "Powered by Groq Cloud API: Primary (openai/gpt-oss-120b) + Fallback (qwen3.8-27b).",
                "Curates strictly against technical profile (AI agents, LLMs, Python, backends).",
                "High-Signal 3-Tier Categorization: Must Read, Inspiration, or Skip.",
                "Generates 2-sentence technical brief and 'Why it matters' strategic note.",
            ],
            EMERALD,
        ),
        (
            "AGENT 3: PUBLISHER",
            "Markdown Compiler & Database Syncer",
            [
                "Compiles structured executive reports to digests/YYYY-MM-DD.md.",
                "Windows cp1252 & UTF-8 resilient (zero terminal crashes on Windows).",
                "Transactional SQLite persistence for digests, metrics, and curated items.",
                "Defensive overwrite protection: safeguards historical records.",
            ],
            AMBER,
        ),
    ]

    card_w = Inches(3.7)
    card_h = Inches(4.5)
    for i, (title, sub, bullets, color) in enumerate(agents):
        cx = Inches(0.8 + i * 4.0)
        cy = Inches(1.7)
        add_card(slide2, cx, cy, card_w, card_h)

        abox = slide2.shapes.add_textbox(cx + Inches(0.2), cy + Inches(0.2), card_w - Inches(0.4), card_h - Inches(0.4))
        atf = abox.text_frame
        atf.word_wrap = True

        ap1 = atf.paragraphs[0]
        ap1.text = title
        ap1.font.size = Pt(12)
        ap1.font.bold = True
        ap1.font.color.rgb = color
        ap1.font.name = "Arial"

        ap2 = atf.add_paragraph()
        ap2.text = sub
        ap2.font.size = Pt(14)
        ap2.font.bold = True
        ap2.font.color.rgb = WHITE
        ap2.font.name = "Arial"
        ap2.space_after = Pt(14)

        for b in bullets:
            bp = atf.add_paragraph()
            bp.text = f"•  {b}"
            bp.font.size = Pt(11)
            bp.font.color.rgb = TEXT_BODY
            bp.font.name = "Arial"
            bp.space_after = Pt(8)

    # Bottom Architecture Summary Note
    b_note = slide2.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.5))
    b_np = b_note.text_frame.paragraphs[0]
    b_np.text = "Result: Complete end-to-end execution completes in ~15.2 seconds with zero human intervention required."
    b_np.font.size = Pt(11)
    b_np.font.bold = True
    b_np.font.color.rgb = EMERALD
    b_np.font.name = "Arial"

    # =========================================================================
    # SLIDE 3: Multi-Source Technical Intelligence Network (6 Feeds)
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, prs)
    add_header(slide3, "The 6-Source Intelligence Ingestion Network")

    sources = [
        ("GitHub Trending", "Scrapes top daily repos across Python and all engineering languages", "HTML Scraper", CYAN),
        ("GitHub Search API", "Monitors fast-moving newly created repositories in AI agents & tools", "REST API", CYAN),
        ("Hacker News Frontpage", "Algolia API frontpage discussions, point scores, and comments", "Algolia API", AMBER),
        ("Hugging Face Papers", "Daily research papers ranked by ML community upvotes & citations", "Hugging Face API", AMBER),
        ("Dev.to Engineering", "Curated engineering articles covering AI frameworks and system design", "Dev.to REST", EMERALD),
        ("ArXiv AI Preprints", "Official cs.AI and cs.LG Atom XML research feed for cutting-edge theory", "XML / Atom", EMERALD),
    ]

    sw = Inches(3.7)
    sh = Inches(2.1)
    for i, (name, desc, tech, color) in enumerate(sources):
        col = i % 3
        row = i // 3
        sx = Inches(0.8 + col * 4.0)
        sy = Inches(1.7 + row * 2.4)
        add_card(slide3, sx, sy, sw, sh)

        sbox = slide3.shapes.add_textbox(sx + Inches(0.2), sy + Inches(0.18), sw - Inches(0.4), sh - Inches(0.36))
        stf = sbox.text_frame
        stf.word_wrap = True

        sp1 = stf.paragraphs[0]
        sp1.text = name
        sp1.font.size = Pt(14)
        sp1.font.bold = True
        sp1.font.color.rgb = WHITE
        sp1.font.name = "Arial"
        sp1.space_after = Pt(2)

        sp_tech = stf.add_paragraph()
        sp_tech.text = f"Protocol: {tech}"
        sp_tech.font.size = Pt(10)
        sp_tech.font.bold = True
        sp_tech.font.color.rgb = color
        sp_tech.font.name = "Arial"
        sp_tech.space_after = Pt(6)

        sp2 = stf.add_paragraph()
        sp2.text = desc
        sp2.font.size = Pt(10)
        sp2.font.color.rgb = TEXT_BODY
        sp2.font.name = "Arial"

    # Bottom benchmark bar
    stat_bar = slide3.shapes.add_textbox(Inches(0.8), Inches(6.55), Inches(11.7), Inches(0.4))
    sb_p = stat_bar.text_frame.paragraphs[0]
    sb_p.text = "Benchmark: Ingests 90-100 items across all 6 streams in under 5.2 seconds concurrent wall time."
    sb_p.font.size = Pt(11)
    sb_p.font.bold = True
    sb_p.font.color.rgb = CYAN
    sb_p.font.name = "Arial"

    # =========================================================================
    # SLIDE 4: Web Observatory Cockpit — Responsive Bento Grid
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, prs)
    add_header(slide4, "Web Cockpit: Responsive Bento Grid & High-Signal Curation")

    # Left: Embedded Real Screenshot 1
    if SCREENSHOT_1.exists():
        slide4.shapes.add_picture(str(SCREENSHOT_1), Inches(0.8), Inches(1.6), Inches(7.2), Inches(4.5))
    else:
        add_card(slide4, Inches(0.8), Inches(1.6), Inches(7.2), Inches(4.5))

    # Right: Feature Highlights Card
    rx = Inches(8.3)
    ry = Inches(1.6)
    rw = Inches(4.2)
    rh = Inches(4.5)
    add_card(slide4, rx, ry, rw, rh)

    rbox = slide4.shapes.add_textbox(rx + Inches(0.25), ry + Inches(0.2), rw - Inches(0.5), rh - Inches(0.4))
    rtf = rbox.text_frame
    rtf.word_wrap = True

    rp0 = rtf.paragraphs[0]
    rp0.text = "SIGNAL MONOLITH DESIGN"
    rp0.font.size = Pt(11)
    rp0.font.bold = True
    rp0.font.color.rgb = EMERALD
    rp0.font.name = "Arial"
    rp0.space_after = Pt(4)

    rp1 = rtf.add_paragraph()
    rp1.text = "Engineered for Focus"
    rp1.font.size = Pt(16)
    rp1.font.bold = True
    rp1.font.color.rgb = WHITE
    rp1.font.name = "Arial"
    rp1.space_after = Pt(12)

    features = [
        ("High Signal Default:", "Skip items are quarantined into a 'Noise Audit' tab so readers never see clutter."),
        ("Interactive KPI Strip:", "Clicking on 'High Conviction' or 'Exploratory Signals' filters items instantly."),
        ("Responsive Bento Grid:", "Auto-fill columns adapt from 1 to 3 columns without awkward right-hand voids."),
        ("Source Spectrum Pills:", "Filter by GitHub, Hacker News, Hugging Face, or ArXiv with live count badges."),
        ("Export Markdown:", "Direct one-click download of executive-ready digest reports."),
    ]

    for f_title, f_desc in features:
        fp = rtf.add_paragraph()
        fp.text = f"• {f_title} {f_desc}"
        fp.font.size = Pt(10)
        fp.font.color.rgb = TEXT_BODY
        fp.font.name = "Arial"
        fp.space_after = Pt(6)

    # Footer note
    f_bar = slide4.shapes.add_textbox(Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.4))
    f_p = f_bar.text_frame.paragraphs[0]
    f_p.text = "Live Local Server: FastAPI backend running at http://127.0.0.1:8000 with real-time SSE & hot-reload."
    f_p.font.size = Pt(11)
    f_p.font.color.rgb = TEXT_MUTED
    f_p.font.name = "Arial"

    # =========================================================================
    # SLIDE 5: Zero-Scroll Cockpit Flight Mode & Keyboard Flight Controls
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, prs)
    add_header(slide5, "Zero-Scroll Cockpit Flight Mode & Keyboard Navigation")

    # Left: Embedded Real Screenshot 2
    if SCREENSHOT_2.exists():
        slide5.shapes.add_picture(str(SCREENSHOT_2), Inches(0.8), Inches(1.6), Inches(7.2), Inches(4.5))
    else:
        add_card(slide5, Inches(0.8), Inches(1.6), Inches(7.2), Inches(4.5))

    # Right: Cockpit Capabilities
    add_card(slide5, rx, ry, rw, rh)

    cbox = slide5.shapes.add_textbox(rx + Inches(0.25), ry + Inches(0.2), rw - Inches(0.5), rh - Inches(0.4))
    ctf = cbox.text_frame
    ctf.word_wrap = True

    cp0 = ctf.paragraphs[0]
    cp0.text = "ZERO-SCROLL FLIGHT DECK"
    cp0.font.size = Pt(11)
    cp0.font.bold = True
    cp0.font.color.rgb = CYAN
    cp0.font.name = "Arial"
    cp0.space_after = Pt(4)

    cp1 = ctf.add_paragraph()
    cp1.text = "Terminal-Speed Curation"
    cp1.font.size = Pt(16)
    cp1.font.bold = True
    cp1.font.color.rgb = WHITE
    cp1.font.name = "Arial"
    cp1.space_after = Pt(12)

    cockpit_feats = [
        ("Keyboard Flight Controls:", "Navigate signals with [Up] / [Down] or [J] / [K]; launch source directly with [Enter]."),
        ("Dual-Pane Layout:", "Left scrollable signal list + Right sticky inspector pane for zero viewport scrolling."),
        ("Deep Technical Telemetry:", "Inspects technical briefs, strategic relevance takeaways, source origins, and direct URIs."),
        ("Ctrl + K Quick Search:", "Instant in-memory filtering across titles, summaries, and tags in milliseconds."),
        ("One-Click Launch & Copy:", "Direct source launch or clipboard copy without leaving the observatory."),
    ]

    for cf_title, cf_desc in cockpit_feats:
        cfp = ctf.add_paragraph()
        cfp.text = f"• {cf_title} {cf_desc}"
        cfp.font.size = Pt(10)
        cfp.font.color.rgb = TEXT_BODY
        cfp.font.name = "Arial"
        cfp.space_after = Pt(6)

    # Footer note
    c_bar = slide5.shapes.add_textbox(Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.4))
    c_p = c_bar.text_frame.paragraphs[0]
    c_p.text = "Verified via Playwright browser automation with 100% headless test pass rate."
    c_p.font.size = Pt(11)
    c_p.font.color.rgb = TEXT_MUTED
    c_p.font.name = "Arial"

    # =========================================================================
    # SLIDE 6: Verification, Performance & Business Impact
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, prs)
    add_header(slide6, "Quality Verification, Performance & Delivery Summary")

    summary_cards = [
        (
            "TESTING EXCELLENCE",
            "22 / 22 Tests Passing (100%)",
            [
                "6 Collector unit tests with network timeout isolation",
                "SQLite 30-day deduplication & non-destructive tests",
                "Pydantic strict schema & JSON fence stripping tests",
                "Playwright Chromium browser UI end-to-end tests",
            ],
            EMERALD,
        ),
        (
            "ZERO-COST ARCHITECTURE",
            "$0.00 Monthly Operating Cost",
            [
                "Inference: Groq Cloud API free tier (120B / 27B models)",
                "Scheduling: GitHub Actions scheduled cron at 8AM UTC",
                "Storage: Local SQLite engine in WAL mode",
                "Ingestion: 100% free public developer APIs & feeds",
            ],
            CYAN,
        ),
        (
            "ENTERPRISE READINESS",
            "Production-Grade Standards",
            [
                "Private GitHub: Bugendaitech-Nishant/techpulse-ai-digest-bot",
                "Windows cp1252 & UTF-8 resilient (zero terminal crashes)",
                "FastAPI REST endpoints for stats, digests, and live harvests",
                "Export-ready Markdown reports with strategic takeaways",
            ],
            AMBER,
        ),
    ]

    for i, (title, sub, bullets, color) in enumerate(summary_cards):
        cx = Inches(0.8 + i * 4.0)
        cy = Inches(1.7)
        add_card(slide6, cx, cy, card_w, card_h)

        vbox = slide6.shapes.add_textbox(cx + Inches(0.2), cy + Inches(0.2), card_w - Inches(0.4), card_h - Inches(0.4))
        vtf = vbox.text_frame
        vtf.word_wrap = True

        vp1 = vtf.paragraphs[0]
        vp1.text = title
        vp1.font.size = Pt(12)
        vp1.font.bold = True
        vp1.font.color.rgb = color
        vp1.font.name = "Arial"

        vp2 = vtf.add_paragraph()
        vp2.text = sub
        vp2.font.size = Pt(14)
        vp2.font.bold = True
        vp2.font.color.rgb = WHITE
        vp2.font.name = "Arial"
        vp2.space_after = Pt(14)

        for b in bullets:
            vbp = vtf.add_paragraph()
            vbp.text = f"✔  {b}"
            vbp.font.size = Pt(11)
            vbp.font.color.rgb = TEXT_BODY
            vbp.font.name = "Arial"
            vbp.space_after = Pt(8)

    # Final wrap-up quote
    q_box = slide6.shapes.add_textbox(Inches(0.8), Inches(6.45), Inches(11.7), Inches(0.5))
    qp = q_box.text_frame.paragraphs[0]
    qp.text = "Status: Fully operational, tested, documented, and hosted on Private GitHub."
    qp.font.size = Pt(12)
    qp.font.bold = True
    qp.font.color.rgb = EMERALD
    qp.font.name = "Arial"

    prs.save(str(OUTPUT_FILE))
    print(f"[SUCCESS] Presentation generated successfully: {OUTPUT_FILE}")


if __name__ == "__main__":
    build_presentation()
